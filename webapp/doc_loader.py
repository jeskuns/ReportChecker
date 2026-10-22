"""
文档加载模块（Web 版）
支持 PDF / Word(.docx) / 图片，统一输出页面图像（PNG bytes）与文本内容
"""

import base64
import io
import os
from typing import List, Optional

import fitz  # PyMuPDF


class LoadedDocument:
    """已加载的文档：统一为若干「页」，每页可提供图像与文本"""

    def __init__(self, doc_id: str, file_path: str):
        self.doc_id = doc_id
        self.file_path = file_path
        self.doc_type = ''           # pdf_text / pdf_image / word / image
        self.page_count = 0
        self._fitz_doc: Optional[fitz.Document] = None
        self._word_text = ''
        self._image_bytes: Optional[bytes] = None
        self.is_text_based = False

    # ---------- 打开 ----------

    def open(self) -> bool:
        ext = os.path.splitext(self.file_path)[1].lower()
        if ext == '.pdf':
            return self._open_pdf()
        if ext in ('.docx', '.doc'):
            return self._open_word()
        if ext in ('.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff', '.webp'):
            return self._open_image()
        raise ValueError(f'不支持的文件格式: {ext}')

    def _open_pdf(self) -> bool:
        self._fitz_doc = fitz.open(self.file_path)
        self.page_count = self._fitz_doc.page_count
        # 判定文字型 / 图片型：前几页文本量
        text_len = 0
        for i in range(min(3, self.page_count)):
            text_len += len(self._fitz_doc[i].get_text().strip())
        self.is_text_based = text_len > 50
        self.doc_type = 'pdf_text' if self.is_text_based else 'pdf_image'
        return True

    def _open_word(self) -> bool:
        from docx import Document
        doc = Document(self.file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(' | '.join(cells))
        self._word_text = '\n'.join(parts)
        self.page_count = 1
        self.doc_type = 'word'
        self.is_text_based = True
        return True

    def _open_image(self) -> bool:
        with open(self.file_path, 'rb') as f:
            self._image_bytes = f.read()
        self.page_count = 1
        self.doc_type = 'image'
        return True

    def close(self):
        if self._fitz_doc is not None:
            try:
                self._fitz_doc.close()
            except Exception:
                pass
            self._fitz_doc = None

    # ---------- 页面内容 ----------

    def page_text(self, page_num: int) -> str:
        """提取页面文本（文字型 PDF / Word 有效）"""
        if self.doc_type == 'word':
            return self._word_text
        if self._fitz_doc is not None and 0 <= page_num < self.page_count:
            return self._fitz_doc[page_num].get_text()
        return ''

    def all_text(self) -> str:
        if self.doc_type == 'word':
            return self._word_text
        if self._fitz_doc is not None:
            return '\n'.join(self._fitz_doc[i].get_text() for i in range(self.page_count))
        return ''

    def page_png(self, page_num: int, dpi: int = 200) -> Optional[bytes]:
        """渲染指定页为 PNG 字节流"""
        if self.doc_type == 'image':
            if page_num != 0:
                return None
            # 统一转 PNG，方便浏览器显示与 AI 识别
            try:
                pix = fitz.Pixmap(self._image_bytes)
                if pix.alpha or pix.colorspace.n > 3:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                return pix.tobytes('png')
            except Exception:
                return self._image_bytes
        if self._fitz_doc is not None and 0 <= page_num < self.page_count:
            page = self._fitz_doc[page_num]
            zoom = dpi / 72.0
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            return pix.tobytes('png')
        return None

    def page_base64(self, page_num: int, dpi: int = 200) -> Optional[str]:
        png = self.page_png(page_num, dpi)
        if png is None:
            return None
        return base64.b64encode(png).decode('ascii')
